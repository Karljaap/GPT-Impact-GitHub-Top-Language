# -*- coding: utf-8 -*-
"""docx_to_tex.py — Convert Tesis_Entrega_version_final.docx to Tesis.tex"""

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from lxml import etree
import re, os

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(BASE_DIR, "Tesis_Entrega_version_final.docx")
OUT  = os.path.join(BASE_DIR, "Tesis.tex")

# Map figure labels (Word caption text) → figure file names
FIGURE_FILES = {
    "Figura 1": "language_distribution_gpt_available_2020_2023.png",
    "Figura 2": "language_distribution_no_gpt_2020_2023.png",
    "Figura 3": "language_trend_gpt_countries_2020_2023.png",
    "Figura 4": "language_trend_no_gpt_countries_2020_2023.png",
    "Figura 5": "TypeScriptsdid_trends12.png",
    "Figura 6": "JavaScriptsdid_trends12.png",
    "Figura 7": "Pythonsdid_trends12.png",
    "Figura 8": "Javasdid_trends12.png",
    "Figura 9": "Gosdid_trends12.png",
    "Figura 10": "Rubysdid_trends12.png",
    "Figura 11": "PHPsdid_trends12.png",
    "Figura 12": "C_hashtagsdid_trends12.png",
    "Figura 13": "C_plussdid_trends12.png",
    "Figura 14": "Csdid_trends12.png",
}

doc = Document(DOCX)

# ── helpers ────────────────────────────────────────────────────────────────────

def esc(s):
    replacements = [
        ("\\",  "BKSL"),
        ("&",   r"\&"),
        ("%",   r"\%"),
        ("$",   r"\$"),
        ("#",   r"\#"),
        ("_",   r"\_"),
        ("{",   r"\{"),
        ("}",   r"\}"),
        ("~",   r"\textasciitilde{}"),
        ("^",   r"\textasciicircum{}"),
        ("BKSL", r"\textbackslash{}"),
    ]
    for old, new in replacements:
        s = s.replace(old, new)
    return s


def has_equation(para_elem):
    """Return True if paragraph contains Word equation (oMath)."""
    ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    return bool(para_elem.findall(f".//{{{ns}}}oMath"))


def runs_to_tex(para):
    out = []
    for run in para.runs:
        t = esc(run.text)
        if not t.strip():
            # keep whitespace-only runs only if they have no formatting
            if not (run.bold or run.italic):
                out.append(t)
            continue
        if run.bold and run.italic:
            t = r"\textbf{\textit{" + t + "}}"
        elif run.bold:
            t = r"\textbf{" + t + "}"
        elif run.italic:
            t = r"\textit{" + t + "}"
        out.append(t)
    return "".join(out)


def table_to_tex(tbl):
    rows = []
    for row in tbl.rows:
        cells = []
        for cell in row.cells:
            txt = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            cells.append(esc(txt))
        rows.append(cells)

    if not rows:
        return ""

    ncols = max(len(r) for r in rows)
    col_spec = "l" + "c" * (ncols - 1) if ncols > 1 else "l"

    lines = [r"\begin{tabular}{" + col_spec + "}", r"\toprule"]
    for i, row in enumerate(rows):
        while len(row) < ncols:
            row.append("")
        lines.append(" & ".join(row) + r" \\")
        if i == 0:
            lines.append(r"\midrule")
    lines += [r"\bottomrule", r"\end{tabular}"]
    return "\n".join(lines)


def is_figure_panel_table(tbl):
    """Return True if the table is a 1x2 or 2x3 figure panel (contains only panel labels)."""
    texts = []
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
    # Panel labels match "(a)", "(b)", "Tendencias", "Pesos"
    panel_words = {"tendencias", "pesos", "(a)", "(b)", "(c)", "(d)", "(e)", "(f)"}
    if texts and all(any(w in t.lower() for w in panel_words) for t in texts):
        return True
    return False


# ── walk document body ─────────────────────────────────────────────────────────

body    = doc.element.body
parts   = []
in_list = False

# Skip all content before the first Heading 1 (title-page paragraphs)
first_section_found = False

# Track last figure caption to know which figure file to use
pending_fig_label = None   # e.g. "Figura 5"

for child in body:
    tag = child.tag.split("}")[-1]

    if tag == "p":
        para  = Paragraph(child, doc)
        style = para.style.name
        line  = runs_to_tex(para)
        raw   = para.text.strip()

        # close open list if style changes
        if in_list and style != "List Paragraph":
            parts.append(r"\end{itemize}")
            in_list = False

        # gate: skip everything before first Heading 1
        if not first_section_found and style != "Heading 1":
            continue

        # skip empty paragraphs
        if not raw:
            # equation placeholder if oMath present but no visible text
            if has_equation(child):
                parts.append(r"\[ \text{[ecuación]} \]")
            else:
                parts.append("")
            continue

        # detect "Índice de Figuras" / "Índice de Tablas" paragraphs
        if raw in ("Índice de Figuras", "Indice de Figuras"):
            parts.append(r"\listoffigures")
            parts.append(r"\newpage")
            continue
        if raw in ("Índice de Tablas", "Indice de Tablas"):
            parts.append(r"\listoftables")
            parts.append(r"\newpage")
            continue

        # equation
        if has_equation(child):
            parts.append(r"\[ \text{[ecuación]} \]")
            continue

        if style == "Heading 1":
            first_section_found = True
            pending_fig_label = None
            parts.append("\n" + r"\section{" + line + "}")

        elif style == "Heading 2":
            parts.append("\n" + r"\subsection{" + line + "}")

        elif style == "List Paragraph":
            if not in_list:
                parts.append(r"\begin{itemize}")
                in_list = True
            parts.append(r"  \item " + line)

        elif style == "Caption":
            # detect figure number
            m = re.match(r"(Figura\s+\d+)", raw, re.IGNORECASE)
            if m:
                pending_fig_label = m.group(1).strip()
                # strip "Figura N" from the caption text to use as caption
                cap_text = raw[len(pending_fig_label):].strip()
                # emit figure environment
                fname = FIGURE_FILES.get(pending_fig_label, "")
                parts.append(r"\begin{figure}[htbp]")
                parts.append(r"\centering")
                if fname:
                    parts.append(r"\includegraphics[width=\textwidth]{" + fname + "}")
                else:
                    parts.append(r"% [imagen no disponible]")
                if cap_text:
                    parts.append(r"\caption{" + esc(cap_text) + "}")
                else:
                    parts.append(r"\caption{" + esc(raw) + "}")
                parts.append(r"\label{fig:" + pending_fig_label.replace(" ", "").lower() + "}")
                parts.append(r"\end{figure}")
            else:
                # table caption or annexe
                parts.append(r"\textbf{" + line + r"}" + "\\\\")

        elif style == "table of figures":
            # skip — replaced by \listoffigures / \listoftables
            pass

        else:
            # detect note lines after figures
            if raw.startswith("Nota.") or raw.startswith("Nota "):
                parts.append(r"\smallskip\noindent\small\textit{" + line + r"}\normalsize")
            else:
                parts.append(line + "\n")

    elif tag == "tbl":
        tbl = Table(child, doc)
        if in_list:
            parts.append(r"\end{itemize}")
            in_list = False

        # skip figure-panel placeholder tables
        if is_figure_panel_table(tbl):
            continue

        parts.append("")
        parts.append(r"\begin{table}[htbp]")
        parts.append(r"\centering")
        parts.append(table_to_tex(tbl))
        parts.append(r"\end{table}")
        parts.append("")

if in_list:
    parts.append(r"\end{itemize}")

# ── post-process: collapse >2 consecutive blank lines ──────────────────────────
cleaned = []
blank_count = 0
for line in parts:
    if line.strip() == "":
        blank_count += 1
        if blank_count <= 2:
            cleaned.append(line)
    else:
        blank_count = 0
        cleaned.append(line)

body_tex = "\n".join(cleaned)

# ── full document ──────────────────────────────────────────────────────────────

PREAMBLE = r"""\documentclass[12pt,a4paper]{article}

% ── encoding & language ──────────────────────────────────────────────────────
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage[spanish]{babel}

% ── typography ───────────────────────────────────────────────────────────────
\usepackage{lmodern}
\usepackage{microtype}
\usepackage{setspace}
\onehalfspacing

% ── page layout ──────────────────────────────────────────────────────────────
\usepackage[top=2.5cm,bottom=2.5cm,left=3cm,right=2.5cm]{geometry}

% ── math ─────────────────────────────────────────────────────────────────────
\usepackage{amsmath,amssymb}

% ── tables ───────────────────────────────────────────────────────────────────
\usepackage{booktabs}
\usepackage{longtable}
\usepackage{array}
\usepackage{multirow}
\usepackage{makecell}
\usepackage{caption}

% ── figures ──────────────────────────────────────────────────────────────────
\usepackage{graphicx}
\graphicspath{{output/figures/}}

% ── references ───────────────────────────────────────────────────────────────
\usepackage{natbib}
\bibliographystyle{apa}

% ── hyperlinks ───────────────────────────────────────────────────────────────
\usepackage[hidelinks]{hyperref}

% ── misc ─────────────────────────────────────────────────────────────────────
\usepackage{parskip}

% ─────────────────────────────────────────────────────────────────────────────
\begin{document}

% ── portada ──────────────────────────────────────────────────────────────────
\begin{titlepage}
\centering
\vspace*{1.5cm}
{\Large \textbf{PONTIFICIA UNIVERSIDAD CATÓLICA DEL PERÚ}}\par
\vspace{0.4cm}
{\large FACULTAD DE CIENCIAS SOCIALES}\par
\vspace{3cm}
{\LARGE \textbf{Impacto de la disponibilidad de ChatGPT en el desarrollo\\[0.3em]
de software orientado a Ciencia de Datos entre 2020 y 2023}}\par
\vspace{2.5cm}
{\large Tesis para obtener el título profesional de\\
Licenciado en Economía presentado por:}\par
\vspace{0.8cm}
{\large Nicho Rosado, Jesus Alberto\\
Janampa Aparicio, Karl Willem}\par
\vspace{1.5cm}
{\large Asesor(es):}\par
{\large Quispe Rojas, Alexander Wilder}\par
\vfill
{\large Lima, 2026}
\end{titlepage}

\newpage
\tableofcontents
\newpage

"""

CLOSING = r"""

\end{document}
"""

full = PREAMBLE + body_tex + CLOSING

with open(OUT, "w", encoding="utf-8") as f:
    f.write(full)

print(f"Saved: {OUT}")
print(f"Lines: {len(full.splitlines())}")
